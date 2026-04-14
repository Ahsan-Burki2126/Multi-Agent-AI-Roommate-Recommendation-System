"""
Seed Production Database with Survey Users
==========================================
Loads the 327 real IUB survey responses into the Vercel/Neon PostgreSQL database.

Usage:
    cd roommate-matching-system

    # Set your Neon DATABASE_URL then run:
    DATABASE_URL="postgresql://user:pass@host/dbname?sslmode=require" python seed_prod.py

    # To wipe existing survey users first:
    DATABASE_URL="..." python seed_prod.py --clear

The script reuses all mapping helpers from load_survey_data.py so behaviour is
identical to local seeding — only the target database changes.
"""

import sys
import os
import argparse
import logging

logging.getLogger("sqlalchemy").setLevel(logging.ERROR)
logging.getLogger("sqlalchemy.engine").setLevel(logging.ERROR)

# ── resolve project root ──────────────────────────────────────────────────────
ROOT = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, ROOT)
os.chdir(ROOT)

from dotenv import load_dotenv
load_dotenv()

# ── validate DATABASE_URL before touching anything ────────────────────────────
DB_URL = os.environ.get("DATABASE_URL", "")
if not DB_URL:
    print(
        "\nERROR: DATABASE_URL is not set.\n"
        "Export your Neon/production connection string first:\n\n"
        '  export DATABASE_URL="postgresql://user:pass@host/dbname?sslmode=require"\n'
        "  python seed_prod.py\n"
    )
    sys.exit(1)

# Force production environment so Flask picks up PostgreSQL config
os.environ["FLASK_ENV"] = "production"

# ── import mapping helpers from the existing loader ───────────────────────────
from load_survey_data import (          # noqa: E402
    map_budget, map_age, map_age_range, map_gender,
    map_location, map_cleanliness, map_schedule,
    map_smoking_ok, map_pets_ok, map_noise_tolerance,
    map_room_type, safe_email, build_bio,
)

import openpyxl                         # noqa: E402


def seed(excel_path: str, clear_existing: bool = False) -> None:
    """Create Flask app with production config and seed survey users."""
    from backend.app import create_app
    from backend.database import db
    from backend.models.user import User
    from backend.models.preference import UserPreference

    print(f"Connecting to production database …")
    app = create_app("production")

    with app.app_context():
        db.create_all()
        print("Tables verified / created.")

        if clear_existing:
            print("Clearing existing survey users …")
            survey_users = User.query.filter(
                User.email.like("%@survey.iub.edu.pk")
            ).all()
            for u in survey_users:
                db.session.delete(u)
            db.session.commit()
            print(f"  Deleted {len(survey_users)} old survey users.")

        wb = openpyxl.load_workbook(excel_path)
        ws = wb.active

        COL = {
            "timestamp": 1, "name": 2, "age": 3, "gender": 4, "budget": 5,
            "location": 6, "room_type": 7, "smoking": 8, "drinking": 9,
            "sharing_willingness": 10, "pets": 11, "cleanliness": 12,
            "sleep_schedule": 13, "noise_tolerance": 14, "study_habits": 15,
            "guest_frequency": 16, "cooking_frequency": 17, "temperament": 18,
            "communication": 19, "introvert_extrovert": 20,
            "cultural_comfort": 21, "same_background_importance": 22,
            "importance_financial": 23, "importance_sleep": 24,
            "importance_cleaning": 25, "importance_communication": 26,
            "hobbies": 27, "describe_yourself": 28, "ideal_roommate": 29,
            "red_flags": 30, "same_field_pref": 31,
            "col32": 32, "col33": 33, "col34": 34,
            "cultural_background": 35, "responsibility": 36, "email": 37,
        }

        def cell(row, key):
            return ws.cell(row, COL[key]).value

        loaded = skipped = errors = 0
        total_rows = ws.max_row - 1
        print(f"Processing {total_rows} survey responses …\n")

        for row_idx in range(2, ws.max_row + 1):
            try:
                raw_name   = cell(row_idx, "name")
                raw_age    = cell(row_idx, "age")
                raw_gender = cell(row_idx, "gender")
                raw_budget = cell(row_idx, "budget")
                raw_email  = cell(row_idx, "email")

                if not raw_name and not raw_email:
                    skipped += 1
                    continue

                survey_email = str(raw_email or "").strip().lower()
                email = survey_email if (survey_email and "@" in survey_email) \
                    else safe_email(raw_name, row_idx)

                if User.query.filter_by(email=email).first():
                    skipped += 1
                    continue

                row_data = {k: cell(row_idx, k) for k in COL}

                user = User()
                user.email     = email
                user.full_name = str(raw_name or "Anonymous").strip() or "Anonymous"
                user.gender    = map_gender(raw_gender)
                user.city      = map_location(cell(row_idx, "location"))
                user.bio       = build_bio(row_data)
                user.is_active = True
                user.set_password("Survey@2024")

                db.session.add(user)
                db.session.flush()

                bmin, bmax = map_budget(raw_budget)
                amin, amax = map_age_range(raw_age)

                pref = UserPreference()
                pref.user_id            = user.user_id
                pref.budget_min         = bmin
                pref.budget_max         = bmax
                pref.preferred_location = map_location(cell(row_idx, "location"))
                pref.gender_preference  = "Any"
                pref.age_min            = amin
                pref.age_max            = amax
                pref.cleanliness_level  = map_cleanliness(cell(row_idx, "cleanliness"))
                pref.schedule           = map_schedule(cell(row_idx, "sleep_schedule"))
                pref.smoking_ok         = map_smoking_ok(cell(row_idx, "smoking"))
                pref.pets_ok            = map_pets_ok(cell(row_idx, "pets"))
                pref.noise_tolerance    = map_noise_tolerance(cell(row_idx, "noise_tolerance"))
                pref.preferred_room_type = map_room_type(cell(row_idx, "room_type"))
                pref.lease_duration_months = 6

                db.session.add(pref)
                loaded += 1

                if loaded % 50 == 0:
                    db.session.commit()
                    print(f"  {loaded} users saved …")

            except Exception as exc:
                errors += 1
                print(f"  [ERROR] Row {row_idx}: {exc}")
                db.session.rollback()

        db.session.commit()

        total_users = User.query.count()
        total_prefs = UserPreference.query.count()

        print(f"\n{'='*50}")
        print(f"  Seeding complete!")
        print(f"  Loaded  : {loaded} new users")
        print(f"  Skipped : {skipped} (already exist or empty rows)")
        print(f"  Errors  : {errors}")
        print(f"  DB total users       : {total_users}")
        print(f"  DB total preferences : {total_prefs}")
        print(f"{'='*50}\n")
        print("Survey users can log in with password: Survey@2024")


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Seed production DB with survey users.")
    parser.add_argument(
        "--excel",
        default="data/AI Roommate Matching Survey (Responses).xlsx",
        help="Path to the survey Excel file",
    )
    parser.add_argument(
        "--clear",
        action="store_true",
        help="Delete existing survey users before seeding",
    )
    args = parser.parse_args()

    if not os.path.exists(args.excel):
        print(f"ERROR: Excel file not found: {args.excel}")
        sys.exit(1)

    seed(args.excel, clear_existing=args.clear)
