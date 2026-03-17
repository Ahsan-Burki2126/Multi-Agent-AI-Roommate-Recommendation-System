"""
Survey Data Loader
==================
Imports the 327 real survey responses from the Excel file into the database.
Maps survey columns to User and UserPreference models.

Usage:
    cd roommate-matching-system
    python load_survey_data.py
"""

import sys
import os
import re
import json
import math
import random
import string
import logging
from datetime import datetime

# Suppress verbose SQLAlchemy logs BEFORE any imports (avoids Unicode console errors on Windows)
logging.getLogger('sqlalchemy').setLevel(logging.ERROR)
logging.getLogger('sqlalchemy.engine').setLevel(logging.ERROR)

# Setup path
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
os.chdir(os.path.dirname(os.path.abspath(__file__)))

# Load .env
from dotenv import load_dotenv
load_dotenv()

import openpyxl
import bcrypt

# --- Mapping helpers -----------------------------------------------------------

def map_budget(budget_text):
    """Map survey budget text to numeric min/max (PKR/month)."""
    mapping = {
        'Low (Economy)':  (3000,  8000),
        'Medium':         (8000,  20000),
        'High':           (20000, 40000),
        'Very High':      (40000, 80000),
        'Flexible':       (5000,  50000),
    }
    return mapping.get(str(budget_text).strip(), (5000, 20000))


def map_age(age_text):
    """Parse age range string like '21–23' → midpoint integer."""
    if not age_text:
        return 20
    text = str(age_text).replace('–', '-').replace('—', '-')
    match = re.search(r'(\d+)\s*[-–]\s*(\d+)', text)
    if match:
        lo, hi = int(match.group(1)), int(match.group(2))
        return (lo + hi) // 2
    match = re.search(r'(\d+)', text)
    if match:
        return int(match.group(1))
    return 20


def map_age_range(age_text):
    """Return (min, max) from age range string."""
    if not age_text:
        return (18, 30)
    text = str(age_text).replace('–', '-').replace('—', '-')
    match = re.search(r'(\d+)\s*[-–]\s*(\d+)', text)
    if match:
        lo, hi = int(match.group(1)), int(match.group(2))
        return (max(18, lo - 2), min(40, hi + 2))
    mid = map_age(age_text)
    return (max(18, mid - 3), min(40, mid + 3))


def map_gender(gender_text):
    """Map survey gender to M/F/Other."""
    g = str(gender_text or '').strip().lower()
    if g == 'male':
        return 'M'
    if g == 'female':
        return 'F'
    return 'Other'


def map_cleanliness(value):
    """Map 1-5 numeric cleanliness to enum."""
    mapping = {
        1: 'Relaxed',
        2: 'Relaxed',
        3: 'Average',
        4: 'Clean',
        5: 'Very Clean',
    }
    try:
        return mapping.get(int(float(value)), 'Average')
    except (TypeError, ValueError):
        return 'Average'


def map_schedule(sleep_text):
    """Map sleep schedule to schedule enum."""
    s = str(sleep_text or '').lower()
    if 'night owl' in s or 'after 12' in s:
        return 'Night Shift'
    if 'early' in s or 'before 11' in s:
        return '9-5 Job'
    if 'student' in s:
        return 'Student'
    if 'varies' in s or 'no strict' in s:
        return 'Flexible'
    return 'Student'


def map_noise_tolerance(noise_text):
    """Map noise tolerance text to 1-10 integer."""
    mapping = {
        'Very low':  2,
        'Low':       4,
        'Medium':    6,
        'High':      8,
        'Very high': 10,
    }
    return mapping.get(str(noise_text or '').strip(), 5)


def map_smoking_ok(smoking_text):
    """Whether the person accepts smoking (for their preference field)."""
    s = str(smoking_text or '').strip().lower()
    return s in ('yes', 'occasionally')


def map_pets_ok(pets_text):
    """Whether the person accepts pets."""
    s = str(pets_text or '').strip().lower()
    return 'no pets' not in s and 'allergic' not in s


def map_room_type(room_text):
    """Map room type preference to enum."""
    s = str(room_text or '').strip().lower()
    if 'single' in s:
        return 'Single'
    if 'shared' in s:
        return 'Shared'
    return 'Any'


def map_location(location_text):
    """Return a clean location string."""
    s = str(location_text or '').strip()
    # Shorten very long descriptions
    if len(s) > 100:
        s = s[:100]
    return s or 'Near University/College Area'


def hash_password(plain):
    """Bcrypt-hash a password."""
    return bcrypt.hashpw(plain.encode('utf-8'), bcrypt.gensalt(rounds=10)).decode('utf-8')


def safe_email(name, idx):
    """Generate a valid email from a name + index."""
    base = re.sub(r'[^a-zA-Z0-9]', '.', str(name or 'user').lower()).strip('.')
    if not base:
        base = 'user'
    return f"{base}.{idx:04d}@survey.iub.edu.pk"


def build_bio(row_data):
    """Build a bio string from free-text survey fields."""
    parts = []
    desc = str(row_data.get('describe_yourself', '') or '').strip()
    if desc and desc.lower() not in ('none', 'n/a', ''):
        parts.append(desc)
    hobbies = str(row_data.get('hobbies', '') or '').strip()
    if hobbies and hobbies.lower() not in ('none', 'n/a', ''):
        parts.append(f"Hobbies: {hobbies}")
    red_flags = str(row_data.get('red_flags', '') or '').strip()
    if red_flags and red_flags.lower() not in ('none', 'n/a', ''):
        parts.append(f"Note: {red_flags}")
    return ' | '.join(parts) if parts else 'Survey participant'


def build_extra_prefs(row_data):
    """Build a JSON dict for extra survey fields not in the core schema."""
    return {
        'drinking':           str(row_data.get('drinking', '') or ''),
        'sharing_willingness': str(row_data.get('sharing_willingness', '') or ''),
        'study_habits':       str(row_data.get('study_habits', '') or ''),
        'guest_frequency':    str(row_data.get('guest_frequency', '') or ''),
        'cooking_frequency':  str(row_data.get('cooking_frequency', '') or ''),
        'temperament':        str(row_data.get('temperament', '') or ''),
        'communication':      str(row_data.get('communication', '') or ''),
        'introvert_extrovert': str(row_data.get('introvert_extrovert', '') or ''),
        'cultural_comfort':   str(row_data.get('cultural_comfort', '') or ''),
        'same_field_pref':    str(row_data.get('same_field_pref', '') or ''),
        'cultural_background': str(row_data.get('cultural_background', '') or ''),
        'responsibility':     str(row_data.get('responsibility', '') or ''),
        'ideal_roommate':     str(row_data.get('ideal_roommate', '') or ''),
        'hobbies':            str(row_data.get('hobbies', '') or ''),
        'importance_financial': str(row_data.get('importance_financial', '') or ''),
        'importance_sleep':   str(row_data.get('importance_sleep', '') or ''),
        'importance_cleaning': str(row_data.get('importance_cleaning', '') or ''),
        'importance_communication': str(row_data.get('importance_communication', '') or ''),
    }


# --- Main loader --------------------------------------------------------------

def load_survey_data(excel_path, clear_existing=False):
    # Patch DevelopmentConfig to disable SQL echo before importing app
    from backend.config import DevelopmentConfig
    DevelopmentConfig.SQLALCHEMY_ECHO = False

    from backend.app import create_app
    from backend.database import db
    from backend.models.user import User
    from backend.models.preference import UserPreference

    app = create_app('development')

    with app.app_context():
        db.create_all()

        if clear_existing:
            print("Clearing existing survey users...")
            # Only delete survey-generated users (emails end with survey.iub.edu.pk)
            survey_users = User.query.filter(
                User.email.like('%@survey.iub.edu.pk')
            ).all()
            for u in survey_users:
                db.session.delete(u)
            db.session.commit()
            print(f"  Deleted {len(survey_users)} old survey users")

        wb = openpyxl.load_workbook(excel_path)
        ws = wb.active

        # Column index map (1-based from survey)
        COL = {
            'timestamp':            1,
            'name':                 2,
            'age':                  3,
            'gender':               4,
            'budget':               5,
            'location':             6,
            'room_type':            7,
            'smoking':              8,
            'drinking':             9,
            'sharing_willingness': 10,
            'pets':                11,
            'cleanliness':         12,
            'sleep_schedule':      13,
            'noise_tolerance':     14,
            'study_habits':        15,
            'guest_frequency':     16,
            'cooking_frequency':   17,
            'temperament':         18,
            'communication':       19,
            'introvert_extrovert': 20,
            'cultural_comfort':    21,
            'same_background_importance': 22,
            'importance_financial': 23,
            'importance_sleep':    24,
            'importance_cleaning': 25,
            'importance_communication': 26,
            'hobbies':             27,
            'describe_yourself':   28,
            'ideal_roommate':      29,
            'red_flags':           30,
            'same_field_pref':     31,
            'col32':               32,
            'col33':               33,
            'col34':               34,
            'cultural_background': 35,
            'responsibility':      36,
            'email':               37,
        }

        def cell(row, key):
            return ws.cell(row, COL[key]).value

        loaded = 0
        skipped = 0
        errors = 0

        print(f"Loading {ws.max_row - 1} survey responses...")

        for row_idx in range(2, ws.max_row + 1):
            try:
                raw_name    = cell(row_idx, 'name')
                raw_age     = cell(row_idx, 'age')
                raw_gender  = cell(row_idx, 'gender')
                raw_budget  = cell(row_idx, 'budget')
                raw_email   = cell(row_idx, 'email')

                # Skip empty rows
                if not raw_name and not raw_email:
                    skipped += 1
                    continue

                # Build email (use survey email if valid, else generate one)
                survey_email = str(raw_email or '').strip().lower()
                if survey_email and '@' in survey_email:
                    email = survey_email
                else:
                    email = safe_email(raw_name, row_idx)

                # Skip if already loaded
                if User.query.filter_by(email=email).first():
                    skipped += 1
                    continue

                # Collect all row data
                row_data = {k: cell(row_idx, k) for k in COL}

                # --- Create User ---
                user = User()
                user.email = email
                user.full_name = str(raw_name or 'Anonymous').strip() or 'Anonymous'
                user.gender = map_gender(raw_gender)
                user.city = map_location(cell(row_idx, 'location'))
                user.bio = build_bio(row_data)
                user.is_active = True
                user.set_password('Survey@2024')   # default password for all survey users

                db.session.add(user)
                db.session.flush()  # get user_id

                # --- Create Preferences ---
                bmin, bmax = map_budget(raw_budget)
                amin, amax = map_age_range(raw_age)

                pref = UserPreference()
                pref.user_id         = user.user_id
                pref.budget_min      = bmin
                pref.budget_max      = bmax
                pref.preferred_location = map_location(cell(row_idx, 'location'))
                pref.gender_preference  = 'Any'
                pref.age_min         = amin
                pref.age_max         = amax
                pref.cleanliness_level = map_cleanliness(cell(row_idx, 'cleanliness'))
                pref.schedule        = map_schedule(cell(row_idx, 'sleep_schedule'))
                pref.smoking_ok      = map_smoking_ok(cell(row_idx, 'smoking'))
                pref.pets_ok         = map_pets_ok(cell(row_idx, 'pets'))
                pref.noise_tolerance = map_noise_tolerance(cell(row_idx, 'noise_tolerance'))
                pref.preferred_room_type = map_room_type(cell(row_idx, 'room_type'))
                pref.lease_duration_months = 6

                db.session.add(pref)
                loaded += 1

                if loaded % 50 == 0:
                    db.session.commit()
                    print(f"  {loaded} users loaded...")

            except Exception as e:
                errors += 1
                print(f"  [ERROR] Row {row_idx}: {e}")
                db.session.rollback()

        db.session.commit()

        total_users = User.query.count()
        total_prefs = UserPreference.query.count()
        print(f"\nDone!")
        print(f"  Loaded:  {loaded} new users")
        print(f"  Skipped: {skipped} (already exist or empty)")
        print(f"  Errors:  {errors}")
        print(f"  DB total users:       {total_users}")
        print(f"  DB total preferences: {total_prefs}")


if __name__ == '__main__':
    import argparse
    parser = argparse.ArgumentParser(description='Load survey data into the database')
    parser.add_argument('--excel', default='data/AI Roommate Matching Survey (Responses).xlsx',
                        help='Path to the Excel file')
    parser.add_argument('--clear', action='store_true',
                        help='Clear existing survey users before loading')
    args = parser.parse_args()

    if not os.path.exists(args.excel):
        print(f"ERROR: Excel file not found: {args.excel}")
        sys.exit(1)

    load_survey_data(args.excel, clear_existing=args.clear)
